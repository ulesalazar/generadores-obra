"""Libro de Excel de la estimación: carátula, resumen, generadores e importación.

Todo lo calculable lleva fórmula: si se corrige una dimensión en Generadores,
se actualizan la cantidad del periodo, el acumulado, los importes y la carátula.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from .modelo import Estimacion

AZUL = "1F3A5F"
NARANJA = "E8772E"
GRIS = "F2F4F7"
ROJO = "F8D7DA"
MONEDA = "#,##0.00"
CANT = "#,##0.000"

_tit = Font(bold=True, size=14, color=AZUL)
_enc = Font(bold=True, color="FFFFFF")
_rel_enc = PatternFill("solid", fgColor=AZUL)
_rel_tot = PatternFill("solid", fgColor=GRIS)
_neg = Font(bold=True)
_linea = Side(style="thin", color="8A94A6")


def _encabezado(ws, fila: int, textos: list[str]) -> None:
    for i, t in enumerate(textos, start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = _enc
        c.fill = _rel_enc
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _anchos(ws, anchos: list[float]) -> None:
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a


def _datos_obra(ws, est: Estimacion, titulo: str) -> int:
    ws["A1"] = titulo
    ws["A1"].font = _tit
    datos = [
        ("Obra", est.obra),
        ("Contrato", est.contrato),
        ("Contratista", est.contratista),
        ("Dependencia", est.dependencia),
        ("Estimación", str(est.numero)),
        ("Periodo", f"del {est.periodo_del} al {est.periodo_al}" if est.periodo_del else est.periodo_al),
    ]
    fila = 2
    for etiqueta, valor in datos:
        if not valor:
            continue
        # Etiqueta y valor en una sola celda para que no dependa del ancho de la columna A.
        ws.cell(row=fila, column=1, value=f"{etiqueta}: {valor}")
        fila += 1
    return fila + 1


def _firmas(ws, fila: int, est: Estimacion, columnas: tuple[int, ...], ancho: int = 2) -> None:
    """Bloque de firmas: una línea por firma, en las columnas dadas y en filas sucesivas."""
    if not est.firmas:
        return
    fila += 2
    for i, (cargo, nombre) in enumerate(est.firmas.items()):
        col = columnas[i % len(columnas)]
        f = fila + 4 * (i // len(columnas))
        for c in range(col, col + ancho):
            ws.cell(row=f, column=c).border = Border(top=_linea)
        ws.cell(row=f, column=col, value=nombre).font = _neg
        ws.cell(row=f + 1, column=col, value=cargo)


def _imprimir(ws, horizontal: bool, filas_titulo: str | None = None) -> None:
    ws.page_setup.orientation = "landscape" if horizontal else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if filas_titulo:
        ws.print_title_rows = filas_titulo


def _nombre(wb: Workbook, nombre: str, ref: str) -> None:
    wb.defined_names[nombre] = DefinedName(nombre, attr_text=ref)


def exportar_excel(est: Estimacion, ruta: str | Path) -> Path:
    ruta = Path(ruta)
    wb = Workbook()

    # --- Generadores (primero, para conocer las celdas de subtotal) ---------
    wg = wb.active
    wg.title = "Generadores"
    _anchos(wg, [10, 9, 8, 10, 14, 46, 8, 10, 10, 10, 13])
    fila = _datos_obra(wg, est, "NÚMEROS GENERADORES")
    celda_periodo: dict[str, str] = {}
    for clave, gen in est.generadores.items():
        con = est.catalogo.get(clave)
        wg.cell(row=fila, column=1, value=clave).font = Font(bold=True, color=NARANJA, size=12)
        wg.cell(row=fila, column=1).alignment = Alignment(vertical="top")
        wg.cell(row=fila, column=2, value=con.descripcion if con else "(no está en el catálogo)").font = _neg
        wg.cell(row=fila, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        wg.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=9)
        wg.row_dimensions[fila].height = 32
        wg.cell(row=fila, column=10, value="Unidad:")
        wg.cell(row=fila, column=11, value=con.unidad if con else "")
        fila += 1
        _encabezado(
            wg,
            fila,
            [
                "No.",
                "Nivel",
                "Eje",
                "Tramo",
                "Croquis / ref.",
                "Descripción",
                "Cant.",
                "Largo",
                "Ancho",
                "Alto",
                "Parcial",
            ],
        )
        primera = fila + 1
        for n, r in enumerate(gen.renglones, start=1):
            fila += 1
            nota = r.nota
            if r.referencia and r.referencia not in nota:
                nota = f"{nota} ({r.referencia})" if nota else r.referencia
            if r.origen:
                nota = f"{nota} [{r.origen}]"
            valores = [n, r.nivel, r.eje, r.tramo, r.croquis, nota, r.cant, r.largo, r.ancho, r.alto]
            for col, v in enumerate(valores, start=1):
                c = wg.cell(row=fila, column=col, value=v if v != "" else None)
                if col >= 7:
                    c.number_format = CANT
            wg.cell(row=fila, column=6).alignment = Alignment(wrap_text=True, vertical="top")
            parcial = wg.cell(
                row=fila,
                column=11,
                value=(
                    f'=G{fila}*IF(H{fila}="",1,H{fila})*IF(I{fila}="",1,I{fila})*IF(J{fila}="",1,J{fila})'
                ),
            )
            parcial.number_format = CANT
        fila += 1
        wg.cell(row=fila, column=10, value="Subtotal").font = _neg
        sub = wg.cell(row=fila, column=11, value=f"=SUM(K{primera}:K{fila - 1})" if gen.renglones else 0)
        sub.number_format = CANT
        sub.font = _neg
        sub.fill = _rel_tot
        for col in range(1, 12):
            wg.cell(row=fila, column=col).border = Border(top=_linea)
        celda_periodo[clave] = f"Generadores!$K${fila}"
        if gen.desperdicio:
            fila += 1
            wg.cell(
                row=fila, column=6, value="Informativo: el desperdicio va en la matriz del precio unitario."
            )
            wg.cell(row=fila, column=6).font = Font(italic=True, color="666666")
            wg.cell(row=fila, column=9, value="Desperdicio")
            wg.cell(row=fila, column=10, value=gen.desperdicio).number_format = "0.0%"
            wg.cell(row=fila, column=11, value=f"=K{fila - 1}*(1+J{fila})").number_format = CANT
            wg.cell(row=fila, column=11).font = Font(italic=True, color="666666")
        fila += 3
    _firmas(wg, fila, est, (2, 6, 9))
    _imprimir(wg, horizontal=True)

    # --- Resumen --------------------------------------------------------------
    wr = wb.create_sheet("Resumen", 0)
    _anchos(wr, [10, 12, 46, 8, 12, 13, 12, 12, 12, 12, 9, 15, 15])
    fila = _datos_obra(wr, est, "RESUMEN DE LA ESTIMACIÓN")
    _encabezado(
        wr,
        fila,
        [
            "Partida",
            "Clave",
            "Descripción",
            "Unidad",
            "Contratado",
            "P. U.",
            "Anterior",
            "Esta estimación",
            "Acumulado",
            "Por ejecutar",
            "% avance",
            "Importe esta estimación",
            "Importe acumulado",
        ],
    )
    wr.freeze_panes = wr.cell(row=fila + 1, column=4)
    primera = fila + 1
    for linea in est.lineas():
        fila += 1
        c = linea.concepto
        valores = [c.partida, c.clave, c.descripcion, c.unidad, c.cantidad, c.precio_unitario, linea.anterior]
        for col, v in enumerate(valores, start=1):
            wr.cell(row=fila, column=col, value=v)
        wr.cell(row=fila, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        wr.cell(row=fila, column=8, value=f"={celda_periodo[c.clave]}" if c.clave in celda_periodo else 0)
        wr.cell(row=fila, column=9, value=f"=G{fila}+H{fila}")
        wr.cell(row=fila, column=10, value=f"=E{fila}-I{fila}")
        wr.cell(row=fila, column=11, value=f"=IF(E{fila}=0,0,I{fila}/E{fila})").number_format = "0.0%"
        wr.cell(row=fila, column=12, value=f"=H{fila}*F{fila}")
        wr.cell(row=fila, column=13, value=f"=I{fila}*F{fila}")
        for col in (5, 7, 8, 9, 10):
            wr.cell(row=fila, column=col).number_format = CANT
        for col in (6, 12, 13):
            wr.cell(row=fila, column=col).number_format = MONEDA
    ultima = fila
    fila += 1
    wr.cell(row=fila, column=11, value="Totales").font = _neg
    for col in (12, 13):
        letra = get_column_letter(col)
        t = wr.cell(row=fila, column=col, value=f"=SUM({letra}{primera}:{letra}{ultima})")
        t.number_format = MONEDA
        t.font = _neg
        t.fill = _rel_tot
    celda_importe = f"Resumen!$L${fila}"
    _firmas(wr, fila, est, (3, 8, 12))
    _imprimir(wr, horizontal=True)
    if ultima >= primera:
        wr.conditional_formatting.add(
            f"I{primera}:J{ultima}",
            FormulaRule(
                formula=[f"$I{primera}>$E{primera}+0.0000001"], fill=PatternFill("solid", fgColor=ROJO)
            ),
        )

    # --- Carátula --------------------------------------------------------------
    wc = wb.create_sheet("Carátula", 0)
    _anchos(wc, [38, 14, 18])
    fila = _datos_obra(wc, est, "CARÁTULA DE ESTIMACIÓN")
    _encabezado(wc, fila, ["Concepto", "Tasa", "Importe"])
    fila += 1
    wc.cell(row=fila, column=1, value="Importe de la estimación (sin IVA)")
    wc.cell(row=fila, column=3, value=f"={celda_importe}").number_format = MONEDA
    f_imp = fila
    fila += 1
    wc.cell(row=fila, column=1, value="Amortización del anticipo")
    wc.cell(row=fila, column=2, value=est.anticipo).number_format = "0.00%"
    wc.cell(row=fila, column=3, value=f"=-C{f_imp}*B{fila}").number_format = MONEDA
    _nombre(wb, "tasa_anticipo", f"'Carátula'!$B${fila}")
    f_amo = fila
    fila += 1
    wc.cell(row=fila, column=1, value="Subtotal").font = _neg
    wc.cell(row=fila, column=3, value=f"=C{f_imp}+C{f_amo}").number_format = MONEDA
    f_sub = fila
    fila += 1
    wc.cell(row=fila, column=1, value="IVA")
    wc.cell(row=fila, column=2, value=est.iva).number_format = "0.00%"
    wc.cell(row=fila, column=3, value=f"=C{f_sub}*B{fila}").number_format = MONEDA
    f_iva = fila
    fila += 1
    wc.cell(row=fila, column=1, value="Total con IVA").font = _neg
    wc.cell(row=fila, column=3, value=f"=C{f_sub}+C{f_iva}").number_format = MONEDA
    f_tot = fila
    filas_ret = []
    for nombre, tasa in est.retenciones.items():
        fila += 1
        wc.cell(row=fila, column=1, value=f"Retención: {nombre.replace('_', ' ')}")
        wc.cell(row=fila, column=2, value=tasa).number_format = "0.000%"
        wc.cell(row=fila, column=3, value=f"=-C{f_imp}*B{fila}").number_format = MONEDA
        filas_ret.append(fila)
    fila += 1
    wc.cell(row=fila, column=1, value="Líquido a pagar").font = Font(bold=True, color=NARANJA, size=12)
    suma_ret = "".join(f"+C{f}" for f in filas_ret)
    liq = wc.cell(row=fila, column=3, value=f"=C{f_tot}{suma_ret}")
    liq.number_format = MONEDA
    liq.font = Font(bold=True, color=NARANJA, size=12)
    liq.fill = _rel_tot
    for f in range(f_imp, fila + 1):
        wc.cell(row=f, column=1).border = Border(bottom=Side(style="hair", color="C9CED6"))
    wc.cell(
        row=fila + 2, column=1, value="Las retenciones se calculan sobre el importe de la estimación sin IVA."
    ).font = Font(italic=True, color="666666")
    _firmas(wc, fila + 3, est, (1, 3), ancho=1)
    _imprimir(wc, horizontal=False)

    # --- Catálogo para importar ----------------------------------------------
    wi = wb.create_sheet("Catalogo_Importar")
    _anchos(wi, [12, 14, 60, 8, 14])
    _imprimir(wi, horizontal=True)
    _encabezado(wi, 1, ["Partida", "Clave", "Descripción", "Unidad", "Cantidad"])
    for f, linea in enumerate(est.lineas(), start=2):
        c = linea.concepto
        for col, v in enumerate([c.partida, c.clave, c.descripcion, c.unidad, round(linea.esta, 4)], start=1):
            wi.cell(row=f, column=col, value=v)

    # --- Advertencias -----------------------------------------------------------
    avisos = est.advertencias()
    wa = wb.create_sheet("Advertencias")
    _anchos(wa, [110])
    wa["A1"] = "ADVERTENCIAS" if avisos else "Sin advertencias"
    wa["A1"].font = _tit
    for f, a in enumerate(avisos, start=3):
        wa.cell(row=f, column=1, value=a)

    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta
